"""Build a concise Vietnamese Lab 3 report from measured pipeline outputs."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import librosa
import numpy as np
import soundfile as sf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


LABELS = ("mo", "dong", "len", "xuong", "dung")
DISPLAY_LABELS = {
    "mo": "mở",
    "dong": "đóng",
    "len": "lên",
    "xuong": "xuống",
    "dung": "dừng",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _accuracy(rows: list[dict[str, str]]) -> float | None:
    if not rows:
        return None
    correct = sum(row.get("true_label") == row.get("prediction") for row in rows)
    return correct / len(rows)


def _add_page_title(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(title)


def _add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    font_size: int = 8,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)


def _add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def _add_image(document: Document, path: Path, width: float = 5.7) -> None:
    if path.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(width))


def _label_name(label: str) -> str:
    return DISPLAY_LABELS.get(label, label)


def _audio_summary(folder: Path) -> dict[str, float] | None:
    paths = sorted(folder.rglob("*.wav")) if folder.is_dir() else []
    if not paths:
        return None
    durations: list[float] = []
    trimmed_durations: list[float] = []
    rms_values: list[float] = []
    for path in paths:
        waveform, sample_rate = sf.read(path, dtype="float32")
        if waveform.ndim == 2:
            waveform = waveform.mean(axis=1)
        durations.append(len(waveform) / sample_rate)
        trimmed, _ = librosa.effects.trim(waveform, top_db=30)
        trimmed_durations.append(len(trimmed) / sample_rate)
        rms_values.append(float(np.sqrt(np.mean(np.square(waveform)))))
    return {
        "count": float(len(paths)),
        "duration_min": min(durations),
        "duration_max": max(durations),
        "trimmed_median": float(np.median(trimmed_durations)),
        "rms_median": float(np.median(rms_values)),
    }


def generate_report(data_dir: Path, results_dir: Path, output_path: Path) -> Path:
    summary_path = results_dir / "summary.json"
    experiment_path = results_dir / "experiment_metrics.csv"
    if not summary_path.is_file() or not experiment_path.is_file():
        raise FileNotFoundError("Run Lab03.py first to generate summary.json and experiment_metrics.csv.")

    summary: dict[str, Any] = json.loads(summary_path.read_text(encoding="utf-8"))
    experiments = read_csv(experiment_path)
    clean_rows = read_csv(results_dir / "likelihood_scores.csv")
    robustness_metrics = read_csv(results_dir / "robustness_metrics.csv")
    robustness_rows: list[dict[str, str]] = []
    if robustness_metrics:
        robustness_rows = read_csv(
            results_dir / f"robustness_predictions_{robustness_metrics[0]['condition']}.csv"
        )
    comparison = read_csv(results_dir / "comparison_dtw_hmm.csv")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    for style_name in ("Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "Arial"
        document.styles[style_name].font.color.rgb = RGBColor(31, 78, 121)
    document.styles["Title"].font.name = "Arial"
    document.styles["Title"].font.size = Pt(20)

    # Page 1: objective and data.
    _add_page_title(document, "BÁO CÁO LAB 3")
    subtitle = document.add_paragraph("Nhận dạng tiếng nói bằng Gaussian HMM")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    for field in ("Họ và tên: ____________________", "Mã sinh viên: ____________________", "Lớp: ____________________"):
        identity = document.add_paragraph(field)
        identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("1. Mục tiêu và dữ liệu", level=1)
    document.add_paragraph(
        "Xây dựng bộ nhận dạng năm lệnh tiếng Việt bằng whole-word Gaussian HMM trái-sang-phải; "
        "trích đặc trưng MFCC động, dùng likelihood để phân lớp, quan sát đường Viterbi, "
        "so sánh với DTW và đánh giá khả năng tổng quát hóa."
    )
    labels = summary.get("labels", list(LABELS))
    train_counts = summary["hmm"].get("n_train_per_class", {})
    document.add_paragraph(
        f"Dữ liệu chính gồm {sum(train_counts.values())} bản ghi train "
        f"({', '.join(f'{_label_name(label)}: {train_counts.get(label, 0)}' for label in labels)}) "
        f"và {summary.get('n_test', len(clean_rows))} bản ghi test. "
        "Tập train và test được giữ tách riêng. Các file được chuẩn hóa thành mono, 16 kHz, PCM 16-bit."
    )
    _add_bullet(
        document,
        "Robustness: đánh giá trên các bản ghi trong dataset/robustness; kết quả này chỉ đại diện "
        "cho điều kiện thu đã ghi lại, không khẳng định tổng quát cho mọi người nói/thiết bị.",
    )
    document.add_heading("2. Thiết kế thực nghiệm", level=1)
    document.add_paragraph(
        "MFCC 13 chiều, cửa sổ 25 ms, bước nhảy 10 ms; đặc trưng động ghép MFCC, Δ và Δ² "
        "thành 39 chiều. Áp dụng utterance-CMVN. HMM dùng Gaussian một thành phần/state, "
        "covariance diagonal, π bắt đầu ở state 0; chỉ cho self-loop và chuyển state kế tiếp."
    )
    document.add_page_break()

    # Page 2: model and algorithms.
    _add_page_title(document, "MÔ HÌNH VÀ THUẬT TOÁN")
    document.add_heading("3. Cấu hình HMM", level=1)
    document.add_paragraph(
        "Mỗi từ có một HMM 5-state. Với state i < 4, xác suất khởi tạo self-loop là 0,7 "
        "và chuyển tiếp là 0,3; state cuối tự lặp với xác suất 1. HMM được huấn luyện bằng "
        "Baum–Welch (EM) thông qua hmmlearn; ma trận chuyển tiếp được cố định để bảo toàn "
        "topology trái-sang-phải."
    )
    document.add_paragraph("A =")
    for state in range(5):
        values = ["0"] * 5
        if state < 4:
            values[state] = "0,7"
            values[state + 1] = "0,3"
        else:
            values[state] = "1"
        document.add_paragraph("[" + "  ".join(values) + "]")
    document.add_heading("4. Diễn giải thuật toán", level=1)
    _add_bullet(
        document,
        "Forward cộng xác suất của mọi đường trạng thái để tính P(X|λ); Viterbi thay phép cộng bằng max "
        "để tìm đường trạng thái tốt nhất. Hai thuật toán trả lời hai câu hỏi khác nhau.",
    )
    _add_bullet(
        document,
        "Trong Baum–Welch, γt(i) là trách nhiệm/occupancy của state i tại frame t; "
        "ξt(i,j) là trách nhiệm của chuyển tiếp i→j giữa hai frame.",
    )
    _add_bullet(
        document,
        "Delta bổ sung xu hướng biến thiên của phổ trong khi HMM mô hình hóa diễn tiến trạng thái; "
        "hai nguồn thông tin không trùng nhau.",
    )
    _add_bullet(
        document,
        "Diagonal covariance giảm số tham số so với full covariance, phù hợp hơn khi ít mẫu; "
        "đổi lại giả định các chiều đặc trưng độc lập có điều kiện.",
    )
    document.add_heading("5. Viterbi alignment", level=1)
    document.add_paragraph(
        "State path cho biết số frame được gán cho từng state, không phải phiên âm phoneme. "
        "Warping path của DTW căn chỉnh hai chuỗi feature; Viterbi căn chỉnh quan sát với state của HMM."
    )
    _add_image(document, results_dir / "viterbi_1_dong_07.png", width=5.1)
    document.add_page_break()

    # Page 3: clean test and controlled experiments.
    _add_page_title(document, "KẾT QUẢ TRÊN TEST SẠCH")
    document.add_heading("6. So sánh hệ thống", level=1)
    metric_rows: list[list[str]] = []
    for row in comparison:
        metric = row.get("metric", "")
        if metric == "accuracy":
            metric_rows.append(["Accuracy", f"{float(row['DTW']):.1%}", f"{float(row['HMM']):.1%}"])
        elif metric == "training_seconds":
            metric_rows.append(["Train time (s)", f"{float(row['DTW']):.3f}", f"{float(row['HMM']):.3f}"])
        elif metric == "test_seconds_per_trial":
            metric_rows.append(
                [
                    "Scoring time/trial (s)",
                    f"{float(row['DTW']):.3f}",
                    f"{float(row['HMM']):.4f}",
                ]
            )
        elif metric == "top_confusion_pair":
            metric_rows.append(["Nhầm lẫn lớn nhất", row["DTW"], row["HMM"]])
    _add_table(document, ["Chỉ số", "DTW (MFCC 13)", "HMM (39-D)"], metric_rows)
    document.add_paragraph(
        "Thời gian là scoring trên cùng test split, không tính trích đặc trưng. DTW dùng MFCC tĩnh 13-D, "
        "HMM dùng MFCC động 39-D; vì cấu hình feature khác nhau, so sánh accuracy không cô lập riêng "
        "tác động của thuật toán. DTW template không có bước fit tham số nên thời gian train bằng 0."
    )
    document.add_heading("7. Thí nghiệm tham số", level=1)
    selected = {
        "states_3",
        "states_5",
        "states_7",
        "features_static",
        "features_dynamic",
        "covariance_diag",
        "covariance_full",
    }
    rows: list[list[str]] = []
    for row in experiments:
        if row.get("experiment") not in selected:
            continue
        accuracy = f"{float(row['accuracy']):.1%}" if row.get("accuracy") else "Không chạy được"
        description = row["experiment"]
        explanation = row.get("error", "")
        rows.append([description, accuracy, explanation])
    _add_table(document, ["Cấu hình", "Accuracy", "Ghi chú"], rows, font_size=7)
    _add_image(document, results_dir / "confusion_hmm.png", width=3.7)
    document.add_page_break()

    # Page 4: robustness.
    _add_page_title(document, "ROBUSTNESS VÀ PHÂN TÍCH LỖI")
    document.add_heading("8. Người nói/điều kiện mới", level=1)
    if robustness_metrics:
        condition_name = robustness_metrics[0]["condition"]
        for row in robustness_metrics:
            accuracy = float(row["accuracy"])
            document.add_paragraph(
                f"Điều kiện {row['condition']}: HMM đúng {accuracy:.1%} "
                f"({round(accuracy * int(row['n_test']))}/{row['n_test']} mẫu trên "
                "tập robustness)."
            )
        if robustness_rows:
            per_class: list[list[str]] = []
            for label in labels:
                label_rows = [row for row in robustness_rows if row["true_label"] == label]
                correct = sum(row["true_label"] == row["prediction"] for row in label_rows)
                per_class.append(
                    [
                        _label_name(label),
                        str(len(label_rows)),
                        f"{correct}/{len(label_rows)}" if label_rows else "0/0",
                    ]
                )
            _add_table(document, ["Lệnh", "Số mẫu", "Đúng"], per_class)
            wrong_rows = [
                [
                    Path(row["file"]).name,
                    _label_name(row["true_label"]),
                    _label_name(row["prediction"]),
                ]
                for row in robustness_rows
                if row["true_label"] != row["prediction"]
            ]
            if wrong_rows:
                document.add_paragraph("Các trường hợp nhận dạng sai:")
                _add_table(document, ["File", "Nhãn thật", "Dự đoán"], wrong_rows)
        _add_image(document, results_dir / f"confusion_robustness_{condition_name}.png", width=4.4)
        training_audio = _audio_summary(data_dir / "train")
        robustness_audio = _audio_summary(data_dir / "robustness")
        if training_audio and robustness_audio:
            document.add_paragraph(
                f"Kiểm tra đầu vào: train có median RMS={training_audio['rms_median']:.3f} "
                f"và median thời lượng sau trim={training_audio['trimmed_median']:.2f}s; "
                f"robustness có median RMS={robustness_audio['rms_median']:.3f} "
                f"và median sau trim={robustness_audio['trimmed_median']:.2f}s "
                f"(thời lượng file gốc {robustness_audio['duration_min']:.2f}–"
                f"{robustness_audio['duration_max']:.2f}s). Kết quả robustness có thể chịu ảnh hưởng "
                "đồng thời bởi mức âm lượng/độ dài khác biệt; chưa thể quy toàn bộ accuracy giảm "
                "cho riêng người nói mới."
            )
    else:
        document.add_paragraph("Chưa có dữ liệu robustness được đánh giá.")
    document.add_paragraph(
        "Nếu người nói mới được thu trong cùng phiên/thiết bị nhưng điều kiện khác, hãy mô tả đúng "
        "điều kiện đó thay vì kết luận speaker-independent. Mỗi lớp nên có thêm nhiều lượt và nhiều "
        "người nói để giảm bất định của accuracy."
    )
    document.add_heading("9. Cẩn trọng khi diễn giải", level=1)
    clean_accuracy = _accuracy(clean_rows)
    if clean_accuracy is not None and robustness_metrics:
        robust_accuracy = float(robustness_metrics[0]["accuracy"])
        document.add_paragraph(
            f"Accuracy trên test sạch là {clean_accuracy:.1%}; robustness là {robust_accuracy:.1%}, "
            f"chênh {100 * (clean_accuracy - robust_accuracy):.1f} điểm phần trăm. "
            "Tập test sạch chỉ có 10 mẫu và robustness có ít mẫu hơn đáng kể; một dự đoán làm "
            "accuracy thay đổi nhiều điểm phần trăm."
        )
    document.add_paragraph(
        "Không chọn cấu hình cuối bằng accuracy test. Accuracy 100% ở một cấu hình trên test nhỏ "
        "là quan sát thăm dò, không chứng minh cấu hình đó tổng quát tốt hơn."
    )
    document.add_page_break()

    # Page 5: discussion and remaining work.
    _add_page_title(document, "THẢO LUẬN VÀ CÔNG VIỆC CÒN LẠI")
    document.add_heading("10. Trả lời câu hỏi phân tích", level=1)
    _add_bullet(
        document,
        "Ba nguyên nhân Baum–Welch hội tụ vào nghiệm kém: khởi tạo không phù hợp; quá ít dữ liệu "
        "so với số state/chiều feature; covariance suy biến hoặc các lần thu có speaker/channel mismatch.",
    )
    _add_bullet(
        document,
        "Whole-word HMM có thể phân lớp từ đơn trong vocabulary cố định mà không cần lexicon. "
        "Lexicon cần thiết khi ghép phone/subword model để tạo mô hình từ/câu.",
    )
    _add_bullet(
        document,
        "Self-loop cho phép state giữ lâu hơn khi người nói kéo dài âm. Với cùng utterance và cùng T, "
        "không cần chia log-likelihood/frame để xếp hạng model; chuẩn hóa hữu ích khi độ dài khác nhau "
        "hoặc đặt ngưỡng reject.",
    )
    _add_bullet(
        document,
        "Argmax luôn chọn một lớp kể cả với từ ngoài vocabulary; cần hiệu chuẩn ngưỡng normalized "
        "log-likelihood hoặc margin top-1/top-2 trên validation và dữ liệu unknown riêng.",
    )
    document.add_heading("11. Phần chưa thể kết luận/hoàn tất", level=1)
    pending: list[str] = []
    if any(row.get("experiment") == "covariance_full" and row.get("error") for row in experiments):
        pending.append(
            "Full covariance thất bại với dữ liệu hiện có do covariance không xác định dương; "
            "báo cáo là thí nghiệm không chạy được, không coi là accuracy 0."
        )
    if not (data_dir / "commands.txt").is_file() or not (data_dir / "contexts.csv").is_file():
        pending.append(
            "Bigram re-ranking chưa chạy: bổ sung corpus commands.txt và contexts.csv có ngữ cảnh thật "
            "nếu cần nộp E6."
        )
    if not any(line.strip() and not line.lstrip().startswith("#") for line in (data_dir.parent / "lexicon.txt").read_text(encoding="utf-8").splitlines()):
        pending.append(
            "Lexicon chưa có phát âm: điền theo phone inventory/phiên âm do giảng viên cung cấp."
        )
    pending.append(
        "Ghi nguồn dữ liệu, người nói theo mã ẩn danh, thiết bị/khoảng cách/nhiễu; tăng số test độc lập "
        "và so sánh với đúng baseline DTW của Lab 2 nếu có."
    )
    for item in pending:
        _add_bullet(document, item)
    document.add_heading("Kết luận", level=1)
    document.add_paragraph(
        "Pipeline MFCC động–Gaussian HMM đã chạy và được so sánh với DTW; robustness cũng đã được "
        "đánh giá trên dữ liệu mới. Kết quả cho thấy khả năng tổng quát hóa giảm so với test sạch, "
        "nên cần thêm dữ liệu và ghi rõ các giới hạn trước khi kết luận về hiệu năng."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
